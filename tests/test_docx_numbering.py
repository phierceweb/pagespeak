from __future__ import annotations

import pytest
from docx import Document

from pagespeak.backends._docx_numbering import (
    LevelDef,
    WordNumbering,
    build_level_defs,
    format_number,
)
from pagespeak.backends._docx_structured import render_markdown


@pytest.mark.parametrize(
    ("value", "fmt", "expected"),
    [
        (3, "decimal", "3"),
        (3, "decimalZero", "03"),
        (12, "decimalZero", "12"),
        (1, "lowerLetter", "a"),
        (26, "lowerLetter", "z"),
        (27, "lowerLetter", "aa"),
        (28, "upperLetter", "BB"),
        (4, "lowerRoman", "iv"),
        (1994, "upperRoman", "MCMXCIV"),
        (5, "none", ""),
        (7, "ordinalText", "7"),
        (0, "upperRoman", "0"),
    ],
)
def test_format_number(value: int, fmt: str, expected: str) -> None:
    assert format_number(value, fmt) == expected


def _defs(*levels: LevelDef, num_id: int = 1) -> dict[tuple[int, int], LevelDef]:
    return {(num_id, ilvl): level for ilvl, level in enumerate(levels)}


_LEVEL_TEXTS = ("%1.", "%1.%2.", "%1.%2.%3.")


def _decimal(depth: int, num_id: int = 1, list_id: str = "") -> dict[tuple[int, int], LevelDef]:
    levels = (LevelDef("decimal", _LEVEL_TEXTS[i], 1, False, list_id=list_id) for i in range(depth))
    return _defs(*levels, num_id=num_id)


def test_advancing_a_level_restarts_every_deeper_level() -> None:
    numbering = WordNumbering(_decimal(2))
    numbering.advance(1, 0)
    numbering.advance(1, 1)
    numbering.advance(1, 1)
    assert numbering.label(1, 1) == "1.2."
    numbering.advance(1, 0)
    numbering.advance(1, 1)
    assert numbering.label(1, 1) == "2.1."


def test_level_not_yet_used_shows_one_below_its_start() -> None:
    # Word shows `0.1` for a Heading 2 before any Heading 1, `1.0.1` for a skipped level.
    numbering = WordNumbering(_decimal(3))
    numbering.advance(1, 1)
    assert numbering.label(1, 1) == "0.1."
    numbering.advance(1, 0)
    assert numbering.label(1, 0) == "1."
    numbering.advance(1, 2)
    assert numbering.label(1, 2) == "1.0.1."


def test_legal_numbering_shows_every_level_as_decimal() -> None:
    numbering = WordNumbering(
        _defs(
            LevelDef("upperRoman", "%1.", 1, False),
            LevelDef("decimal", "%1.%2", 1, True),
        )
    )
    numbering.advance(1, 0)
    numbering.advance(1, 0)
    numbering.advance(1, 1)
    assert numbering.label(1, 0) == "II."
    assert numbering.label(1, 1) == "2.1"


def test_no_label_for_bullets_or_unknown_lists() -> None:
    numbering = WordNumbering(_defs(LevelDef("bullet", "•", 1, False)))
    numbering.advance(1, 0)
    assert numbering.label(1, 0) == ""
    assert numbering.label(9, 0) == ""


def test_lists_over_one_definition_share_a_count() -> None:
    numbering = WordNumbering(_decimal(1, list_id="5") | _decimal(1, num_id=3, list_id="5"))
    numbering.advance(1, 0)
    numbering.advance(3, 0)
    numbering.advance(1, 0)
    assert numbering.label(1, 0) == "3."


def test_start_override_restarts_its_level_the_first_time_its_list_uses_it() -> None:
    restart = LevelDef("decimal", "%1.", 1, False, list_id="5", start_override=True)
    numbering = WordNumbering(_decimal(1, list_id="5") | {(4, 0): restart})
    numbering.advance(1, 0)
    numbering.advance(1, 0)
    numbering.advance(4, 0)
    assert numbering.label(4, 0) == "1."
    numbering.advance(1, 0)
    assert numbering.label(1, 0) == "2."
    numbering.advance(4, 0)
    assert numbering.label(4, 0) == "3."


_NUMBERING = """
<w:abstractNum w:abstractNumId="7">
  <w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="upperRoman"/>
    <w:lvlText w:val="%1."/><w:pPr><w:ind w:left="720"/></w:pPr></w:lvl>
  <w:lvl w:ilvl="1"><w:numFmt w:val="decimal"/><w:lvlText w:val="%1.%2"/><w:isLgl/></w:lvl>
</w:abstractNum>
<w:num w:numId="4"><w:abstractNumId w:val="7"/>
  <w:lvlOverride w:ilvl="0"><w:startOverride w:val="5"/></w:lvlOverride>
</w:num>
"""


def test_build_level_defs_reads_levels_and_start_override(make_docx) -> None:
    defs = build_level_defs(Document(str(make_docx(document_xml="", numbering_xml=_NUMBERING))))
    assert defs[(4, 0)] == LevelDef("upperRoman", "%1.", 5, False, 720, "7", True)
    # No `w:start`: the spec's default is 0.
    assert defs[(4, 1)] == LevelDef("decimal", "%1.%2", 0, True, None, "7", False)


def test_build_level_defs_without_definitions_is_empty(make_docx) -> None:
    assert build_level_defs(Document(str(make_docx(document_xml="")))) == {}


def test_build_level_defs_reads_each_format(make_docx) -> None:
    numbering = (
        '<w:abstractNum w:abstractNumId="0"><w:lvl w:ilvl="0"><w:numFmt w:val="decimal"/></w:lvl>'
        '<w:lvl w:ilvl="1"><w:numFmt w:val="lowerLetter"/></w:lvl></w:abstractNum>'
        '<w:abstractNum w:abstractNumId="5"><w:lvl w:ilvl="0"><w:numFmt w:val="bullet"/></w:lvl>'
        '</w:abstractNum><w:abstractNum w:abstractNumId="6"><w:lvl w:ilvl="0"/></w:abstractNum>'
        '<w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num>'
        '<w:num w:numId="7"><w:abstractNumId w:val="5"/></w:num>'
        '<w:num w:numId="8"><w:abstractNumId w:val="6"/></w:num>'
    )
    defs = build_level_defs(Document(str(make_docx(document_xml="", numbering_xml=numbering))))
    assert [defs[k].fmt for k in ((1, 0), (1, 1), (7, 0), (8, 0))] == [
        "decimal",
        "lowerLetter",
        "bullet",
        "decimal",
    ]


_LIST_STYLE = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
<w:style w:type="numbering" w:styleId="Bullets"><w:name w:val="Bullets"/>
  <w:pPr><w:numPr><w:numId w:val="2"/></w:numPr></w:pPr></w:style>
</w:styles>"""


def test_build_level_defs_follows_a_list_style(make_docx) -> None:
    # `w:numStyleLink` defers to the definition a numbering style points at.
    numbering = (
        '<w:abstractNum w:abstractNumId="3"><w:numStyleLink w:val="Bullets"/></w:abstractNum>'
        '<w:abstractNum w:abstractNumId="4"><w:styleLink w:val="Bullets"/>'
        '<w:lvl w:ilvl="0"><w:numFmt w:val="bullet"/><w:pPr><w:ind w:left="360"/></w:pPr>'
        "</w:lvl></w:abstractNum>"
        '<w:num w:numId="1"><w:abstractNumId w:val="3"/></w:num>'
        '<w:num w:numId="2"><w:abstractNumId w:val="4"/></w:num>'
    )
    path = make_docx(document_xml="", numbering_xml=numbering, styles_xml=_LIST_STYLE)
    linked = build_level_defs(Document(str(path)))[(1, 0)]
    assert (linked.fmt, linked.left, linked.list_id) == ("bullet", 360, "4")


# Word's numbered-heading template: `Heading N` linked to a multilevel list on the style.
_HEADING_STYLES = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
<w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="heading 1"/>
  <w:pPr><w:numPr><w:numId w:val="1"/></w:numPr></w:pPr></w:style>
<w:style w:type="paragraph" w:styleId="Heading2"><w:name w:val="heading 2"/>
  <w:pPr><w:numPr><w:ilvl w:val="1"/><w:numId w:val="1"/></w:numPr></w:pPr></w:style>
<w:style w:type="paragraph" w:styleId="Heading3"><w:name w:val="heading 3"/>
  <w:pPr><w:numPr><w:ilvl w:val="2"/><w:numId w:val="1"/></w:numPr></w:pPr></w:style>
{extra}
</w:styles>"""


def _lvl(ilvl: int, fmt: str, text: str, start: str = '<w:start w:val="1"/>') -> str:
    return f'<w:lvl w:ilvl="{ilvl}">{start}<w:numFmt w:val="{fmt}"/><w:lvlText w:val="{text}"/></w:lvl>'


def _heading_numbering(*levels: str, more: str = "") -> str:
    return (
        f'<w:abstractNum w:abstractNumId="5">{"".join(levels)}</w:abstractNum>'
        f'<w:num w:numId="1"><w:abstractNumId w:val="5"/></w:num>{more}'
    )


_DECIMAL_HEADINGS = _heading_numbering(*(_lvl(i, "decimal", t) for i, t in enumerate(_LEVEL_TEXTS)))


def _h(level: int, text: str) -> str:
    return (
        f'<w:p><w:pPr><w:pStyle w:val="Heading{level}"/></w:pPr><w:r><w:t>{text}</w:t></w:r></w:p>'
    )


def _body(text: str = "Body.") -> str:
    return f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>"


def _render(make_docx, xml: str, numbering: str = _DECIMAL_HEADINGS, extra: str = "") -> list[str]:
    styles = _HEADING_STYLES.format(extra=extra)
    docx = make_docx(document_xml=xml, numbering_xml=numbering, styles_xml=styles)
    return render_markdown(Document(str(docx)), None).splitlines()


def test_style_numbered_headings_carry_their_word_numbers(make_docx) -> None:
    xml = (
        _h(1, "Overview")
        + _body()
        + _h(2, "Scope")
        + _body()
        + _h(3, "Limits")
        + _body()
        + _h(1, "Installation")
        + _body()
        + _h(2, "Requirements")
        + _body()
    )
    lines = _render(make_docx, xml)
    assert "# 1. Overview" in lines, "\n".join(lines)
    assert "## 1.1. Scope" in lines
    assert "### 1.1.1. Limits" in lines
    assert "# 2. Installation" in lines
    assert "## 2.1. Requirements" in lines


def test_empty_numbered_heading_still_takes_its_number(make_docx) -> None:
    # Word shows a number on an empty numbered paragraph.
    empty = '<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr></w:p>'
    lines = _render(make_docx, _h(1, "First") + _body() + empty + _h(1, "Third") + _body())
    assert "# 1. First" in lines, "\n".join(lines)
    assert "# 3. Third" in lines


def test_heading_number_formats_follow_each_level(make_docx) -> None:
    numbering = _heading_numbering(_lvl(0, "upperRoman", "%1."), _lvl(1, "upperLetter", "%2)"))
    xml = (
        _h(1, "Alpha")
        + _h(2, "First part")
        + _body()
        + _h(2, "Second part")
        + _body()
        + _h(1, "Beta")
        + _body()
    )
    lines = _render(make_docx, xml, numbering)
    assert "# I. Alpha" in lines, "\n".join(lines)
    assert "## A) First part" in lines
    assert "## B) Second part" in lines
    assert "# II. Beta" in lines


def _list_item(num_id: int, text: str) -> str:
    return (
        f'<w:p><w:pPr><w:numPr><w:ilvl w:val="0"/><w:numId w:val="{num_id}"/></w:numPr></w:pPr>'
        f"<w:r><w:t>{text}</w:t></w:r></w:p>"
    )


def test_numbered_paragraph_on_the_heading_list_advances_it(make_docx) -> None:
    xml = _h(1, "One") + _body() + _list_item(1, "On the heading list") + _h(1, "Three") + _body()
    lines = _render(make_docx, xml)
    assert "# 1. One" in lines, "\n".join(lines)
    assert "# 3. Three" in lines


def test_numbered_paragraph_on_a_list_sharing_the_definition_advances_it(make_docx) -> None:
    numbering = _DECIMAL_HEADINGS + '<w:num w:numId="3"><w:abstractNumId w:val="5"/></w:num>'
    xml = _h(1, "One") + _body() + _list_item(3, "Same definition") + _h(1, "Next") + _body()
    lines = _render(make_docx, xml, numbering)
    assert "# 3. Next" in lines, "\n".join(lines)


def test_numbered_paragraph_in_a_table_advances_the_count(make_docx) -> None:
    table = (
        '<w:tbl><w:tblGrid><w:gridCol w:w="5000"/></w:tblGrid>'
        f"<w:tr><w:tc>{_h(1, 'In a table')}</w:tc></w:tr></w:tbl>"
    )
    lines = _render(make_docx, _h(1, "One") + _body() + table + _h(1, "After") + _body())
    assert "# 3. After" in lines, "\n".join(lines)


_TEXT_BOX_NS = (
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
    'xmlns:v="urn:schemas-microsoft-com:vml"'
)


def test_numbered_paragraph_in_a_table_text_box_is_not_counted(make_docx) -> None:
    # A text box is its own story, and Word writes it twice (DrawingML + VML fallback).
    box = f"<w:txbxContent>{_h(1, 'In a text box')}</w:txbxContent>"
    run = (
        f"<w:r><mc:AlternateContent {_TEXT_BOX_NS}>"
        f'<mc:Choice Requires="wps"><w:drawing><wps:txbx>{box}</wps:txbx></w:drawing></mc:Choice>'
        f"<mc:Fallback><w:pict><v:shape><v:textbox>{box}</v:textbox></v:shape></w:pict>"
        "</mc:Fallback></mc:AlternateContent></w:r>"
    )
    table = (
        '<w:tbl><w:tblGrid><w:gridCol w:w="5000"/></w:tblGrid>'
        f"<w:tr><w:tc><w:p>{run}</w:p></w:tc></w:tr></w:tbl>"
    )
    lines = _render(make_docx, _h(1, "One") + _body() + table + _h(1, "After") + _body())
    assert "# 2. After" in lines, "\n".join(lines)


def test_heading_2_before_any_heading_1(make_docx) -> None:
    xml = _h(2, "Preface") + _body() + _h(1, "Intro") + _body() + _h(2, "Scope") + _body()
    lines = _render(make_docx, xml)
    assert "## 0.1. Preface" in lines, "\n".join(lines)
    assert "# 1. Intro" in lines
    assert "## 1.1. Scope" in lines


def test_levels_without_a_start_count_from_zero(make_docx) -> None:
    numbering = _heading_numbering(_lvl(0, "decimal", "%1.", start=""))
    lines = _render(make_docx, _h(1, "One") + _body() + _h(1, "Two") + _body(), numbering)
    assert "# 0. One" in lines, "\n".join(lines)
    assert "# 1. Two" in lines


def test_headings_numbered_through_a_list_style(make_docx) -> None:
    # `w:numStyleLink` defers to the definition a numbering style points at.
    list_style = (
        '<w:style w:type="numbering" w:styleId="HeadList"><w:name w:val="HeadList"/>'
        '<w:pPr><w:numPr><w:numId w:val="2"/></w:numPr></w:pPr></w:style>'
    )
    numbering = (
        '<w:abstractNum w:abstractNumId="5"><w:numStyleLink w:val="HeadList"/></w:abstractNum>'
        '<w:abstractNum w:abstractNumId="6"><w:styleLink w:val="HeadList"/>'
        f"{_lvl(0, 'decimal', '%1.')}{_lvl(1, 'decimal', '%1.%2.')}</w:abstractNum>"
        '<w:num w:numId="1"><w:abstractNumId w:val="5"/></w:num>'
        '<w:num w:numId="2"><w:abstractNumId w:val="6"/></w:num>'
    )
    xml = _h(1, "One") + _body() + _h(2, "Sub") + _body() + _h(1, "Two") + _body()
    lines = _render(make_docx, xml, numbering, extra=list_style)
    assert "# 1. One" in lines, "\n".join(lines)
    assert "## 1.1. Sub" in lines
    assert "# 2. Two" in lines
