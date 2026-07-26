"""Builds the golden `.docx` for the structure-faithful DOCX reader.

Generated rather than committed as a binary so the fixture stays diffable — a
reviewer sees which shape changed, which an opaque zip does not allow.

python-docx supplies the parts that must be genuinely valid (styles.xml so
`para.style.name` resolves to `Heading 1`, numbering.xml so `numPr.ilvl` is
real). The run containers below are injected as raw XML because python-docx
does not emit them when writing.

Content is deliberately generic: the fixture is a carrier for *structure*.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_M = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def _numbered(doc: Document, text: str, ilvl: int, num_id: int = 1) -> Paragraph:
    """A body paragraph carrying an explicit `w:numPr` at `ilvl`.

    The reader reads outline depth from `numPr`, not from indentation, so the
    fixture has to set it explicitly — `add_paragraph(style="List Number")`
    alone does not pin a level.
    """
    p = doc.add_paragraph(text)
    ppr = p._p.get_or_add_pPr()
    num_pr = ppr.makeelement(qn("w:numPr"), {})
    ilvl_el = ppr.makeelement(qn("w:ilvl"), {qn("w:val"): str(ilvl)})
    num_el = ppr.makeelement(qn("w:numId"), {qn("w:val"): str(num_id)})
    num_pr.append(ilvl_el)
    num_pr.append(num_el)
    ppr.append(num_pr)
    return p


def _raw_runs(doc: Document, inner_xml: str, ilvl: int | None = None) -> Paragraph:
    """A paragraph whose run content is hand-authored OOXML."""
    p = doc.add_paragraph()
    if ilvl is not None:
        ppr = p._p.get_or_add_pPr()
        num_pr = ppr.makeelement(qn("w:numPr"), {})
        num_pr.append(ppr.makeelement(qn("w:ilvl"), {qn("w:val"): str(ilvl)}))
        num_pr.append(ppr.makeelement(qn("w:numId"), {qn("w:val"): "1"}))
        ppr.append(num_pr)
    from lxml import etree

    frag = etree.fromstring(f'<w:wrap xmlns:w="{_W}" xmlns:m="{_M}">{inner_xml}</w:wrap>')
    for child in frag:
        p._p.append(child)
    return p


def _r(text: str) -> str:
    return f"<w:r><w:t xml:space='preserve'>{text}</w:t></w:r>"


def build_structured_docx(path: Path) -> Path:
    """Write the golden fixture to `path` and return it."""
    doc = Document()

    doc.add_heading("Widget Maintenance Guide", level=1)
    doc.add_paragraph("Introductory paragraph with enough words to be a real body.")

    doc.add_heading("Safety checks", level=1)
    doc.add_heading("Before you begin", level=2)

    _numbered(doc, "Disconnect the power supply", 0)
    _numbered(doc, "Verify the indicator is dark", 1)
    _numbered(doc, "Wait for the capacitor to discharge", 2)

    # Prose-shaped outline item: `demote_prose_heading` targets exactly this.
    # It is also a BULLET parent (ilvl0 of this numbering is a bullet), so the
    # ordered child below it must restart at 1 — no heading intervenes to mask it.
    _numbered(
        doc,
        "The interlock will not release until the residual charge has fully "
        "dissipated, which can take several minutes in a cold enclosure.",
        0,
    )
    _numbered(doc, "Confirm the interlock has released", 1)

    # A run of bare labels: the empty-shell demote targets exactly this.
    doc.add_heading("Connector types", level=2)
    for label in ("spade", "ferrule", "ring", "blade", "pin"):
        _numbered(doc, label, 1)

    doc.add_heading("Notation", level=2)
    # Sub/superscript split across runs — the shatter shape.
    _raw_runs(
        doc,
        _r("Coolant is CO") + "<w:r><w:rPr><w:vertAlign w:val='subscript'/></w:rPr>"
        "<w:t>2</w:t></w:r>"
        + _r(" at 25 m")
        + "<w:r><w:rPr><w:vertAlign w:val='superscript'/></w:rPr>"
        "<w:t>3</w:t></w:r>" + _r(" per hour."),
    )
    # `w:tab` between runs must not fuse the words either side.
    _raw_runs(doc, _r("Torque") + "<w:r><w:tab/></w:r>" + _r("setting"))

    # Run containers whose text must survive. Tracked-insertion text is accepted
    # text, part of the document.
    doc.add_heading("Containers", level=2)
    _raw_runs(
        doc,
        f"<w:ins w:id='1' w:author='a' w:date='2020-01-01T00:00:00Z'>"
        f"{_r('Inserted clause retained.')}</w:ins>",
    )
    _raw_runs(doc, f"<w:sdt><w:sdtContent>{_r('Content control text.')}</w:sdtContent></w:sdt>")
    _raw_runs(doc, f"<w:fldSimple w:instr='PAGE'>{_r('Field text.')}</w:fldSimple>")
    _raw_runs(doc, f"<w:smartTag w:element='x'>{_r('Smart tag text.')}</w:smartTag>")
    _raw_runs(doc, "<m:oMath><m:r><m:t>E = mc^2</m:t></m:r></m:oMath>")

    doc.save(str(path))
    return path
